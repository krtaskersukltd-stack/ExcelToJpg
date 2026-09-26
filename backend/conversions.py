"""Additional source/target conversion helpers for the Excel converter API."""

from __future__ import annotations

import json
import os
import re
import time
import uuid
from pathlib import Path
from xml.etree import ElementTree as ET

import pandas as pd
from docx import Document
from fastapi import HTTPException
from openpyxl import Workbook
from pypdf import PdfReader

FILE_TO_EXCEL_KINDS = {
    "jpg": {"jpg", "jpeg"},
    "jpeg": {"jpg", "jpeg"},
    "png": {"png"},
    "pdf": {"pdf"},
    "bank_statement_pdf": {"pdf"},
    "csv": {"csv"},
    "tsv": {"tsv", "txt"},
    "json": {"json"},
    "xml": {"xml"},
    "txt": {"txt", "text", "log"},
    "text": {"txt", "text", "log"},
    "notepad": {"txt", "text", "log"},
    "word": {"docx"},
    "docx": {"docx"},
    "ods": {"ods"},
    "vcf": {"vcf", "vcard"},
}

SOURCE_KIND_ALIASES = {
    "bank_statement": "bank_statement_pdf",
    "bankstatement": "bank_statement_pdf",
    "bank_pdf": "bank_statement_pdf",
    "doc": "word",
    "docx": "word",
    "vcard": "vcf",
    "jpeg": "jpeg",
}


def safe_stem(value: str) -> str:
    cleaned = "".join(char if char.isalnum() or char in "-_" else "_" for char in str(value))
    return cleaned.strip("_")[:60] or "sheet"


def escape_xml(value: str) -> str:
    return (
        str(value)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&apos;")
    )


def normalize_source_kind(source_kind: str) -> str:
    kind = source_kind.lower().strip().replace("-", "_").replace(" ", "_")
    return SOURCE_KIND_ALIASES.get(kind, kind)


BANK_DATE_RE = re.compile(
    r"^(?P<date>\d{1,2}[-/\.]\d{1,2}[-/\.]\d{2,4}|\d{1,2}\s+[A-Za-z]{3,9}\s+\d{2,4})\b"
)
BANK_AMOUNT_RE = re.compile(r"(?<![\w.])(-?\d{1,3}(?:,\d{3})*(?:\.\d{1,2})?|-?\d+(?:\.\d{1,2})?)(?![\w.])")


def pdf_rows(file_path: str):
    rows = []
    for page in PdfReader(file_path).pages:
        for line in (page.extract_text() or "").splitlines():
            cells = [cell.strip() for cell in re.split(r"\t+|\s{2,}", line) if cell.strip()]
            if cells:
                rows.append(cells)
    if not rows:
        raise HTTPException(status_code=422, detail="No extractable table text was found in this PDF")
    return rows


def bank_statement_pdf_rows(file_path: str):
    header = ["Date", "Description", "Debit", "Credit", "Balance"]
    transactions = []
    for page in PdfReader(file_path).pages:
        for raw_line in (page.extract_text() or "").splitlines():
            line = " ".join(raw_line.split())
            if not line or len(line) < 6:
                continue
            date_match = BANK_DATE_RE.match(line)
            if not date_match:
                continue
            date_value = date_match.group("date")
            remainder = line[date_match.end():].strip()
            amounts = BANK_AMOUNT_RE.findall(remainder)
            if not amounts:
                continue
            description = remainder
            for amount in reversed(amounts):
                idx = description.rfind(amount)
                if idx >= 0:
                    description = description[:idx].rstrip(" -|/")
            description = description.strip() or "Transaction"
            amounts = [a.replace(",", "") for a in amounts]
            debit = credit = balance = ""
            if len(amounts) >= 3:
                debit, credit, balance = amounts[-3], amounts[-2], amounts[-1]
                try:
                    if float(debit or 0) == 0:
                        debit = ""
                    if float(credit or 0) == 0:
                        credit = ""
                except ValueError:
                    pass
            elif len(amounts) == 2:
                first, balance = amounts
                try:
                    value = float(first)
                except ValueError:
                    value = 0
                if value < 0 or "withdraw" in description.lower() or "debit" in description.lower():
                    debit = first.lstrip("-")
                else:
                    credit = first.lstrip("-")
            else:
                credit = amounts[0]
            transactions.append([date_value, description, debit, credit, balance])
    if transactions:
        return [header] + transactions
    return pdf_rows(file_path)


def word_rows(file_path: str):
    document = Document(file_path)
    rows = []
    for table in document.tables:
        for table_row in table.rows:
            cells = [cell.text.strip() for cell in table_row.cells]
            if any(cells):
                rows.append(cells)
    if not rows:
        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if text:
                cells = [cell.strip() for cell in re.split(r"\t+| {2,}", text) if cell.strip()]
                rows.append(cells if len(cells) > 1 else [text])
    if not rows:
        raise HTTPException(status_code=422, detail="No table or text content found in this Word document")
    return rows


def json_rows(file_path: str):
    with open(file_path, "r", encoding="utf-8-sig") as handle:
        payload = json.load(handle)
    if isinstance(payload, list):
        if not payload:
            raise HTTPException(status_code=422, detail="JSON array is empty")
        if all(isinstance(item, dict) for item in payload):
            dataframe = pd.json_normalize(payload)
        else:
            dataframe = pd.DataFrame({"value": payload})
    elif isinstance(payload, dict):
        if payload and all(isinstance(value, list) for value in payload.values()):
            dataframe = pd.DataFrame(payload)
        else:
            dataframe = pd.json_normalize(payload)
            if len(dataframe.columns) <= 1:
                dataframe = pd.DataFrame(list(payload.items()), columns=["key", "value"])
    else:
        dataframe = pd.DataFrame({"value": [payload]})
    return [list(dataframe.columns)] + dataframe.fillna("").astype(str).values.tolist()


def text_rows(file_path: str):
    with open(file_path, "r", encoding="utf-8-sig", errors="replace") as handle:
        lines = [line.rstrip("\n\r") for line in handle if line.strip()]
    if not lines:
        raise HTTPException(status_code=422, detail="The text file is empty")
    sample = "\n".join(lines[:20])
    if "\t" in sample:
        delimiter = "\t"
    elif sample.count("|") >= len(lines[:20]):
        delimiter = "|"
    elif sample.count(";") >= len(lines[:20]):
        delimiter = ";"
    elif sample.count(",") >= len(lines[:20]):
        delimiter = ","
    else:
        return [["Line", "Text"]] + [[index + 1, line] for index, line in enumerate(lines)]
    return [[cell.strip() for cell in line.split(delimiter)] for line in lines]


def tsv_rows(file_path: str):
    dataframe = pd.read_csv(file_path, sep="\t", dtype=str, engine="python")
    return [list(dataframe.columns)] + dataframe.fillna("").values.tolist()


def xml_rows(file_path: str):
    try:
        dataframe = pd.read_xml(file_path)
        if dataframe is not None and not dataframe.empty:
            return [list(dataframe.columns)] + dataframe.fillna("").astype(str).values.tolist()
    except Exception:
        pass

    tree = ET.parse(file_path)
    root = tree.getroot()
    records = []
    children = list(root)
    if children and all(len(list(child)) > 0 for child in children):
        for child in children:
            record = {re.sub(r"\{.*\}", "", item.tag): (item.text or "").strip() for item in list(child)}
            if child.attrib:
                record.update({f"@{key}": value for key, value in child.attrib.items()})
            if record:
                records.append(record)
    if not records:
        records = [
            {
                re.sub(r"\{.*\}", "", root.tag): (root.text or "").strip(),
                **{f"@{key}": value for key, value in root.attrib.items()},
            }
        ]
    dataframe = pd.DataFrame(records).fillna("")
    if dataframe.empty:
        raise HTTPException(status_code=422, detail="No tabular data found in this XML file")
    return [list(dataframe.columns)] + dataframe.astype(str).values.tolist()


def ods_rows(file_path: str):
    try:
        dataframe = pd.read_excel(file_path, engine="odf")
    except Exception as error:
        raise HTTPException(status_code=422, detail=f"Could not read ODS file: {error}") from error
    return [list(dataframe.columns)] + dataframe.fillna("").astype(str).values.tolist()


def vcf_rows(file_path: str):
    with open(file_path, "r", encoding="utf-8-sig", errors="replace") as handle:
        content = handle.read()
    cards = re.split(r"(?i)BEGIN:VCARD", content)
    contacts = []
    for card in cards:
        if "END:VCARD" not in card.upper():
            continue
        fields = {
            "Full Name": "",
            "Phone": "",
            "Email": "",
            "Organization": "",
            "Title": "",
            "Address": "",
            "Note": "",
        }
        phones, emails = [], []
        for raw_line in card.splitlines():
            line = raw_line.strip()
            if not line or ":" not in line:
                continue
            key, value = line.split(":", 1)
            key_upper = key.upper()
            value = value.strip()
            if key_upper.startswith("FN"):
                fields["Full Name"] = value
            elif key_upper.startswith("N") and not fields["Full Name"]:
                parts = [part for part in value.split(";") if part]
                fields["Full Name"] = " ".join(reversed(parts)) if parts else value
            elif key_upper.startswith("TEL"):
                phones.append(value)
            elif key_upper.startswith("EMAIL"):
                emails.append(value)
            elif key_upper.startswith("ORG"):
                fields["Organization"] = value.replace(";", " ").strip()
            elif key_upper.startswith("TITLE"):
                fields["Title"] = value
            elif key_upper.startswith("ADR"):
                fields["Address"] = ", ".join(part for part in value.split(";") if part)
            elif key_upper.startswith("NOTE"):
                fields["Note"] = value
        fields["Phone"] = "; ".join(phones)
        fields["Email"] = "; ".join(emails)
        if any(fields.values()):
            contacts.append(fields)
    if not contacts:
        raise HTTPException(status_code=422, detail="No contacts found in this VCF file")
    columns = list(contacts[0].keys())
    return [columns] + [[contact.get(column, "") for column in columns] for contact in contacts]


def rows_to_xlsx(rows, source_name: str, output_dir: str) -> str:
    job_id = uuid.uuid4().hex[:12]
    output_name = f"{safe_stem(Path(source_name).stem)}_{job_id}.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Extracted Data"
    for row in rows:
        sheet.append(["" if cell is None else cell for cell in list(row)])
    workbook.save(os.path.join(output_dir, output_name))
    return output_name


def source_kind_to_rows(kind: str, input_path: str, ocr_rows_from_image):
    if kind in {"jpg", "jpeg", "png"}:
        return ocr_rows_from_image(input_path)
    if kind == "pdf":
        return pdf_rows(input_path)
    if kind == "bank_statement_pdf":
        return bank_statement_pdf_rows(input_path)
    if kind == "csv":
        dataframe = pd.read_csv(input_path, dtype=str, engine="python")
        return [list(dataframe.columns)] + dataframe.fillna("").values.tolist()
    if kind == "tsv":
        return tsv_rows(input_path)
    if kind in {"txt", "text", "notepad"}:
        return text_rows(input_path)
    if kind == "json":
        return json_rows(input_path)
    if kind == "xml":
        return xml_rows(input_path)
    if kind in {"word", "docx"}:
        return word_rows(input_path)
    if kind == "ods":
        return ods_rows(input_path)
    if kind == "vcf":
        return vcf_rows(input_path)
    raise HTTPException(status_code=400, detail="Unsupported source converter")


def workbook_sheets_to_json(workbook_sheets, output_dir: str, extract_sheet_data):
    job_id = uuid.uuid4().hex[:12]
    payload = {}
    for sheet_name, dataframe in workbook_sheets:
        payload[sheet_name] = json.loads(dataframe.fillna("").to_json(orient="records", force_ascii=False))
    output_name = f"conversion_{job_id}.json"
    with open(os.path.join(output_dir, output_name), "w", encoding="utf-8") as handle:
        json.dump(payload if len(payload) > 1 else next(iter(payload.values()), []), handle, ensure_ascii=False, indent=2)
    return {"filename": output_name, "type": "json", "sheet_data": extract_sheet_data(workbook_sheets)}


def workbook_sheets_to_tally(workbook_sheets, output_dir: str, extract_sheet_data):
    job_id = uuid.uuid4().hex[:12]
    vouchers = []
    for sheet_name, dataframe in workbook_sheets:
        columns = [str(column).strip().lower() for column in dataframe.columns]

        def pick(*names):
            for name in names:
                if name in columns:
                    return dataframe.columns[columns.index(name)]
            return None

        date_col = pick("date", "voucher date", "txn date", "transaction date")
        narration_col = pick("narration", "description", "particulars", "remarks", "memo")
        debit_col = pick("debit", "dr", "withdrawal", "amount debit")
        credit_col = pick("credit", "cr", "deposit", "amount credit")
        amount_col = pick("amount", "value", "txn amount")
        ledger_col = pick("ledger", "account", "party", "ledger name")
        voucher_col = pick("voucher type", "type", "vch type")

        for index, row in dataframe.iterrows():
            amount = ""
            is_debit = True
            if debit_col is not None and str(row.get(debit_col, "")).strip():
                amount = str(row.get(debit_col, "")).strip()
                is_debit = True
            elif credit_col is not None and str(row.get(credit_col, "")).strip():
                amount = str(row.get(credit_col, "")).strip()
                is_debit = False
            elif amount_col is not None:
                amount = str(row.get(amount_col, "")).strip()
                is_debit = True
            else:
                continue
            amount = re.sub(r"[^\d.\-]", "", amount)
            if not amount:
                continue
            date_raw = str(row.get(date_col, "")).strip() if date_col is not None else ""
            date_digits = re.sub(r"\D", "", date_raw)
            if len(date_digits) == 8:
                tally_date = date_digits
            elif len(date_digits) == 6:
                tally_date = "20" + date_digits
            else:
                tally_date = time.strftime("%Y%m%d")
            narration_value = (
                row.get(narration_col, f"{sheet_name} row {index + 1}")
                if narration_col is not None
                else f"{sheet_name} row {index + 1}"
            )
            narration = escape_xml(narration_value)
            ledger = escape_xml(row.get(ledger_col, "Party") if ledger_col is not None else "Party")
            voucher_type = escape_xml(row.get(voucher_col, "Journal") if voucher_col is not None else "Journal")
            remote_id = f"{job_id}-{safe_stem(sheet_name)}-{index + 1}"
            vouchers.append(
                f"""      <TALLYMESSAGE xmlns:UDF="TallyUDF">
        <VOUCHER REMOTEID="{remote_id}" VCHTYPE="{voucher_type}" ACTION="Create">
          <DATE>{tally_date}</DATE>
          <NARRATION>{narration}</NARRATION>
          <VOUCHERTYPENAME>{voucher_type}</VOUCHERTYPENAME>
          <ALLLEDGERENTRIES.LIST>
            <LEDGERNAME>{ledger}</LEDGERNAME>
            <ISDEEMEDPOSITIVE>{"Yes" if is_debit else "No"}</ISDEEMEDPOSITIVE>
            <AMOUNT>{"-" if is_debit else ""}{escape_xml(amount)}</AMOUNT>
          </ALLLEDGERENTRIES.LIST>
        </VOUCHER>
      </TALLYMESSAGE>"""
            )

    if not vouchers:
        raise HTTPException(
            status_code=422,
            detail="Could not map Excel columns to Tally vouchers. Include Amount/Debit/Credit columns.",
        )

    xml_body = f"""<?xml version="1.0" encoding="UTF-8"?>
<ENVELOPE>
  <HEADER>
    <TALLYREQUEST>Import Data</TALLYREQUEST>
  </HEADER>
  <BODY>
    <IMPORTDATA>
      <REQUESTDESC>
        <REPORTNAME>Vouchers</REPORTNAME>
      </REQUESTDESC>
      <REQUESTDATA>
{chr(10).join(vouchers)}
      </REQUESTDATA>
    </IMPORTDATA>
  </BODY>
</ENVELOPE>
"""
    output_name = f"tally_import_{job_id}.xml"
    with open(os.path.join(output_dir, output_name), "w", encoding="utf-8") as handle:
        handle.write(xml_body)
    return {"filename": output_name, "type": "xml", "sheet_data": extract_sheet_data(workbook_sheets)}


def workbook_to_xls(workbook_sheets, output_dir: str, extract_sheet_data):
    import xlwt

    job_id = uuid.uuid4().hex[:12]
    book = xlwt.Workbook()
    for sheet_name, dataframe in workbook_sheets:
        sheet = book.add_sheet(safe_stem(sheet_name)[:31] or "Sheet1")
        for col_index, column in enumerate(dataframe.columns):
            sheet.write(0, col_index, str(column)[:32767])
        for row_index, row in enumerate(dataframe.itertuples(index=False), start=1):
            if row_index > 65535:
                break
            for col_index, value in enumerate(row):
                cell = "" if pd.isna(value) else str(value)[:32767]
                sheet.write(row_index, col_index, cell)
    output_name = f"conversion_{job_id}.xls"
    book.save(os.path.join(output_dir, output_name))
    return {"filename": output_name, "type": "xls", "sheet_data": extract_sheet_data(workbook_sheets)}


def workbook_to_xlsx(workbook_sheets, output_dir: str, extract_sheet_data):
    job_id = uuid.uuid4().hex[:12]
    output_name = f"conversion_{job_id}.xlsx"
    with pd.ExcelWriter(os.path.join(output_dir, output_name), engine="openpyxl") as writer:
        for sheet_name, dataframe in workbook_sheets:
            dataframe.to_excel(writer, index=False, sheet_name=safe_stem(sheet_name)[:31] or "Sheet1")
    return {"filename": output_name, "type": "xlsx", "sheet_data": extract_sheet_data(workbook_sheets)}
